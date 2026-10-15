"""Generate test fixtures using only the project's existing dependencies.

Creates:
- tests/fixtures/sample.docx  — multi-section doc with headings and a table
- tests/fixtures/sample.pdf   — minimal valid PDF with 2 pages of text

Run: uv run python tests/generate_fixtures.py
"""

import struct
from pathlib import Path


def generate_sample_docx(output_path: Path) -> None:
    """Create a DOCX with headings, paragraphs, and a table."""
    from docx import Document

    doc = Document()
    doc.core_properties.title = "Sample DOCX for Testing"

    doc.add_heading("Sample Document for Testing", level=0)

    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(
        "This is a sample DOCX document used for testing the DOCX parser. "
        "It contains headings at multiple levels, body paragraphs, and a table."
    )

    doc.add_heading("Background", level=2)
    doc.add_paragraph(
        "The document parser extracts text while preserving heading hierarchy. "
        "Each heading starts a new section, and paragraphs under a heading are "
        "grouped together as part of that section."
    )

    doc.add_heading("Data Summary", level=2)
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    headers = ["Metric", "Value", "Status"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    data = [
        ["Documents Processed", "1234", "Active"],
        ["Average Chunk Size", "512 tokens", "Optimal"],
        ["Retrieval Accuracy", "94.5 percent", "Good"],
    ]
    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, cell_text in enumerate(row_data):
            table.rows[row_idx].cells[col_idx].text = cell_text

    doc.add_heading("Technical Details", level=1)
    doc.add_paragraph(
        "The parsing pipeline supports multiple document formats. Each format "
        "has a dedicated parser that extracts text while preserving structural "
        "information like headings and tables."
    )

    doc.add_heading("Conclusion", level=2)
    doc.add_paragraph(
        "This sample document demonstrates the key features that our parser "
        "must handle correctly: heading hierarchy, body paragraphs, and tabular data."
    )

    doc.save(str(output_path))


def generate_sample_pdf(output_path: Path) -> None:
    """Create a minimal 2-page PDF with text content.

    Uses raw PDF syntax — no external library needed beyond the stdlib.
    """
    page1_text = (
        "Sample Document for Testing. "
        "Introduction. "
        "This is a sample PDF document used for testing the PDF parser. "
        "It contains multiple pages with text content for verification. "
        "The parser should extract this text page by page. "
        "Background information about the RAG system and how it processes documents."
    )
    page2_text = (
        "Technical Details. "
        "The parsing pipeline supports PDF and DOCX formats. "
        "Each format has a dedicated parser. "
        "Conclusion. "
        "This sample verifies multi-page extraction works correctly."
    )

    pages_text = [page1_text, page2_text]
    objects: list[bytes] = []
    offsets: list[int] = []
    current_offset = 0

    header = b"%PDF-1.4\n"
    current_offset += len(header)

    # Build content streams and page objects
    content_obj_ids = []
    page_obj_ids = []
    font_obj_id = 0
    pages_obj_id = 0
    catalog_obj_id = 0

    # Object numbering: 1-based
    # 1: Catalog, 2: Pages, 3: Font
    # Then for each page: content stream, page object
    obj_num = 4

    # Pre-assign IDs
    font_obj_id = 3
    pages_obj_id = 2
    catalog_obj_id = 1

    for text in pages_text:
        content_obj_ids.append(obj_num)
        obj_num += 1
        page_obj_ids.append(obj_num)
        obj_num += 1

    total_objects = obj_num - 1

    # Build all objects as bytes
    all_data = bytearray(header)

    def add_obj(obj_id: int, content: bytes) -> None:
        nonlocal current_offset
        offsets.append((obj_id, len(all_data)))
        obj_bytes = f"{obj_id} 0 obj\n".encode() + content + b"\nendobj\n"
        all_data.extend(obj_bytes)

    # Object 1: Catalog
    add_obj(1, f"<< /Type /Catalog /Pages {pages_obj_id} 0 R >>".encode())

    # Object 2: Pages
    kids = " ".join(f"{pid} 0 R" for pid in page_obj_ids)
    add_obj(2, f"<< /Type /Pages /Kids [{kids}] /Count {len(pages_text)} >>".encode())

    # Object 3: Font
    add_obj(3, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    # Content streams and pages
    for i, text in enumerate(pages_text):
        # Wrap text into lines
        words = text.split()
        lines = []
        current_line: list[str] = []
        for word in words:
            current_line.append(word)
            if len(" ".join(current_line)) > 70:
                lines.append(" ".join(current_line))
                current_line = []
        if current_line:
            lines.append(" ".join(current_line))

        # Build stream content
        stream_lines = ["BT", f"/F1 12 Tf", "72 720 Td", "14 TL"]
        for line in lines:
            # Escape special PDF characters
            safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream_lines.append(f"({safe}) '")
        stream_lines.append("ET")
        stream_content = "\n".join(stream_lines).encode()

        # Content stream object
        add_obj(
            content_obj_ids[i],
            f"<< /Length {len(stream_content)} >>\nstream\n".encode()
            + stream_content
            + b"\nendstream",
        )

        # Page object
        add_obj(
            page_obj_ids[i],
            (
                f"<< /Type /Page /Parent {pages_obj_id} 0 R "
                f"/MediaBox [0 0 612 792] "
                f"/Contents {content_obj_ids[i]} 0 R "
                f"/Resources << /Font << /F1 {font_obj_id} 0 R >> >> >>"
            ).encode(),
        )

    # Cross-reference table
    xref_offset = len(all_data)
    offsets_sorted = sorted(offsets, key=lambda x: x[0])

    xref = bytearray()
    xref.extend(b"xref\n")
    xref.extend(f"0 {total_objects + 1}\n".encode())
    xref.extend(b"0000000000 65535 f \n")

    offset_map = {obj_id: off for obj_id, off in offsets_sorted}
    for obj_id in range(1, total_objects + 1):
        off = offset_map.get(obj_id, 0)
        xref.extend(f"{off:010d} 00000 n \n".encode())

    all_data.extend(xref)

    # Trailer
    trailer = (
        f"trailer\n<< /Size {total_objects + 1} /Root {catalog_obj_id} 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    )
    all_data.extend(trailer.encode())

    output_path.write_bytes(bytes(all_data))


if __name__ == "__main__":
    fixtures_dir = Path(__file__).parent / "fixtures"
    fixtures_dir.mkdir(exist_ok=True)

    pdf_path = fixtures_dir / "sample.pdf"
    docx_path = fixtures_dir / "sample.docx"

    generate_sample_pdf(pdf_path)
    print(f"Created {pdf_path}")

    generate_sample_docx(docx_path)
    print(f"Created {docx_path}")

    print("Done! Fixtures ready.")
